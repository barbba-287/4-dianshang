"""运营待办格式化 Excel 导出测试。"""

import hashlib
import io

from openpyxl import load_workbook

from app.exports import operations_todo_xlsx


def test_operations_todo_xlsx_has_formatted_centered_bordered_wrapped_cells():
    content, meta = operations_todo_xlsx([{
        "type": "inventory_alert",
        "priority": "critical",
        "status": "open",
        "title": "库存低于补货点 · DEMO-TEA-GREEN-250",
        "reason": "当前库存 4，安全库存 5，补货点 8",
        "warehouse_id": 1,
    }])
    assert content.startswith(b"PK")
    assert meta["sha256"] == hashlib.sha256(content).hexdigest()
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["今日待办清单"]
    assert [cell.value for cell in sheet[1]] == ["待办类型", "优先级", "状态", "待办事项", "原因", "仓库"]
    assert [cell.value for cell in sheet[2]] == ["库存告警", "紧急", "待处理", "库存低于补货点 · DEMO-TEA-GREEN-250", "当前库存 4，安全库存 5，补货点 8", "1"]
    for row in sheet.iter_rows(min_row=1, max_row=2):
        for cell in row:
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
            assert cell.alignment.wrap_text is True
            assert cell.border.left.style == "thin"
            assert cell.border.right.style == "thin"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F2"
    assert sheet.column_dimensions["D"].width <= 42
    assert sheet.column_dimensions["E"].width <= 42


def test_operations_todo_xlsx_empty_result_keeps_formatted_header():
    content, meta = operations_todo_xlsx([])
    sheet = load_workbook(io.BytesIO(content))["今日待办清单"]
    assert meta["row_count"] == 0
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == ["待办类型", "优先级", "状态", "待办事项", "原因", "仓库"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F1"
    assert sheet["A1"].fill.fgColor.rgb.endswith("1F4E78")
