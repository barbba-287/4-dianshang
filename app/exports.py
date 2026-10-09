"""受控的只读运营数据导出。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OPERATIONS_TODO_COLUMNS: tuple[tuple[str, str], ...] = (
    ("type_label", "待办类型"),
    ("priority_label", "优先级"),
    ("status_label", "状态"),
    ("title", "待办事项"),
    ("reason", "原因"),
    ("warehouse_id", "仓库"),
)

_TODO_TYPE_LABELS = {
    "inventory_alert": "库存告警",
    "replenishment_review": "补货待确认",
    "purchase_draft": "采购草稿",
    "sync_failure": "同步异常",
    "content_review": "内容审核",
    "inbound_confirmation": "入库确认",
}
_TODO_PRIORITY_LABELS = {"critical": "紧急", "high": "高", "medium": "中", "low": "低"}
_TODO_STATUS_LABELS = {
    "open": "待处理",
    "acknowledged": "已确认",
    "resolved": "已解决",
    "suggested": "待处理",
    "draft": "草稿",
    "failed": "失败",
    "partial": "部分成功",
    "stalled": "停滞",
    "received": "已收货待确认",
    "in_review": "审核中",
    "review_required": "待审核",
    "quality_failed": "质量检查失败",
    "changes_requested": "待修改",
    "rejected": "已驳回",
    "approved": "已通过",
}


def _todo_export_row(item: dict) -> dict:
    """Add operator-facing labels without changing raw API enum values."""
    return {
        **item,
        "type_label": _TODO_TYPE_LABELS.get(item.get("type"), item.get("type") or "未知"),
        "priority_label": _TODO_PRIORITY_LABELS.get(item.get("priority"), item.get("priority") or "未知"),
        "status_label": _TODO_STATUS_LABELS.get(item.get("status"), item.get("status") or "未知"),
    }


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _safe_csv_cell(value: object) -> str:
    """Prevent spreadsheet formula execution while retaining visible text."""
    text = _cell(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _style_sheet(sheet, *, header_row: int = 1, freeze: str = "A2") -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="1F2937")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = centered
        cell.border = border
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = centered
            cell.border = border
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[header_row].height = 28


def _auto_width(sheet, *, minimum: int = 10, maximum: int = 42) -> None:
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max((len(str(cell.value or "").replace("\n", " ")) for cell in column_cells), default=minimum)
        sheet.column_dimensions[letter].width = min(max(longest + 2, minimum), maximum)


def fulfillment_proposal_xlsx(orders: list[dict], *, meta: dict) -> tuple[bytes, dict]:
    """Build a compact, operator-facing manual-fulfillment workbook."""
    workbook = Workbook()
    order_sheet = workbook.active
    order_sheet.title = "待发订单清单"
    order_headers = ["平台", "店铺", "订单号", "订单状态", "下单时间", "支付时间", "订单金额", "退款金额", "币种", "数据完整度", "来源", "人工核对"]
    order_sheet.append(order_headers)
    status_labels = {"paid": "已支付", "pending": "待处理", "fulfilled": "已履约事实", "completed": "已完成", "cancelled": "已取消", "closed": "已关闭", "unknown": "未知"}
    completeness_labels = {"complete": "完整", "partial": "部分", "insufficient": "数据不足", "unknown": "数据不足"}
    source_labels = {"mock": "模拟/回放", "fixture": "模拟/回放", "json": "文件回放", "shopify_live": "真实 Shopify", "live": "真实来源"}
    for order in orders:
        checks = []
        if order.get("data_completeness") != "complete": checks.append("数据不完整，需人工核对")
        if order.get("refund_amount") not in (None, 0, Decimal("0")): checks.append("存在退款金额，需回平台核对")
        if order.get("simulated"): checks.append("模拟/回放数据，非真实生产订单")
        order_sheet.append([order.get("platform"), order.get("store_ref"), order.get("external_order_no"), status_labels.get(order.get("order_status"), "未知"), order.get("external_created_at"), order.get("paid_at"), order.get("gross_amount"), order.get("refund_amount"), order.get("currency"), completeness_labels.get(order.get("data_completeness"), "数据不足"), source_labels.get(order.get("source_mode"), "其他来源"), "；".join(checks) or "履约前回平台复核"])
    _style_sheet(order_sheet); _auto_width(order_sheet, maximum=42)

    line_sheet = workbook.create_sheet("待拣商品明细")
    line_sheet.append(["平台", "店铺", "订单号", "商品 SKU", "内部商品", "规格", "待拣数量", "取消数量", "退款数量", "金额", "币种", "映射状态", "数据完整度", "人工核对"])
    mapping_labels = {"mapped": "已映射", "unmapped": "未映射"}
    for order in orders:
        for line in order.get("lines", []):
            net_qty = line.get("ordered_qty", 0) - line.get("cancelled_qty", 0) - line.get("refunded_qty", 0)
            mapped = line.get("internal_sku_id") is not None
            checks = []
            if not mapped: checks.append("需要建立商品映射")
            if line.get("data_completeness") != "complete": checks.append("行数据不完整，需人工核对")
            if order.get("simulated"): checks.append("模拟/回放数据")
            line_sheet.append([order.get("platform"), order.get("store_ref"), order.get("external_order_no"), line.get("external_sku"), line.get("internal_product_title") or ("已映射" if mapped else "未映射"), line.get("internal_variant_label") or "—", net_qty, line.get("cancelled_qty"), line.get("refunded_qty"), line.get("gross_amount"), line.get("currency"), mapping_labels.get(line.get("mapping_status"), "未映射"), completeness_labels.get(line.get("data_completeness"), "数据不足"), "；".join(checks) or "履约前复核"])
    _style_sheet(line_sheet); _auto_width(line_sheet, maximum=32)

    output = io.BytesIO(); workbook.save(output); content = output.getvalue()
    return content, {"order_count": len(orders), "row_count": sum(len(order.get("lines", [])) for order in orders), "sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content), "format": "xlsx"}
def operations_todo_xlsx(items: Iterable[dict]) -> tuple[bytes, dict]:
    """Build a readable, formatted Excel workbook for operators."""
    rows = [_todo_export_row(item) for item in items]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "今日待办清单"
    headers = [label for _key, label in OPERATIONS_TODO_COLUMNS]
    sheet.append(headers)
    _style_sheet(sheet)
    for row in rows:
        sheet.append([_safe_csv_cell(row.get(key)) for key, _label in OPERATIONS_TODO_COLUMNS])
    _style_sheet(sheet)
    _auto_width(sheet)
    for row_index, row in enumerate(rows, start=2):
        longest = max((len(_cell(row.get(key))) for key, _label in OPERATIONS_TODO_COLUMNS), default=0)
        sheet.row_dimensions[row_index].height = min(max(24, 18 + (longest // 32) * 15), 72)
    output = io.BytesIO(); workbook.save(output); content = output.getvalue()
    return content, {"row_count": len(rows), "column_keys": [key for key, _label in OPERATIONS_TODO_COLUMNS], "sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content), "format": "xlsx"}


def operations_todo_csv(items: Iterable[dict]) -> tuple[bytes, dict]:
    """Serialize authorized todo rows to deterministic Excel-friendly CSV."""
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow([label for _key, label in OPERATIONS_TODO_COLUMNS])
    count = 0
    for item in items:
        count += 1
        display_item = _todo_export_row(item)
        writer.writerow([_safe_csv_cell(display_item.get(key)) for key, _label in OPERATIONS_TODO_COLUMNS])
    # utf-8-sig makes the file open correctly in common Chinese Excel setups.
    content = output.getvalue().encode("utf-8-sig")
    return content, {
        "row_count": count,
        "column_keys": [key for key, _label in OPERATIONS_TODO_COLUMNS],
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
        "format": "csv",
    }
